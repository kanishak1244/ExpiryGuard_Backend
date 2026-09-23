"""
build_pptx.py
Generates the complete 12-slide 16:9 widescreen PowerPoint presentation (.pptx)
for Dawaiflow Hackathon Presentation with real embedded screenshots and speaker notes.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_FILE = "Dawaiflow_Hackathon_Presentation.pptx"

# Color Palette Constants
COLOR_BG = RGBColor(15, 23, 42)          # Slate 900
COLOR_CARD = RGBColor(30, 41, 59)        # Slate 800
COLOR_CARD_BORDER = RGBColor(51, 65, 85) # Slate 700
COLOR_EMERALD = RGBColor(16, 185, 129)   # Primary Green
COLOR_EMERALD_LIGHT = RGBColor(52, 211, 153)
COLOR_SKY = RGBColor(56, 189, 248)       # Sky Blue
COLOR_AMBER = RGBColor(251, 191, 36)     # Amber
COLOR_ROSE = RGBColor(244, 63, 94)       # Rose / Red
COLOR_WHITE = RGBColor(248, 250, 252)    # Pure Text
COLOR_SECONDARY = RGBColor(148, 163, 184)# Subtitles / muted
COLOR_MUTED = RGBColor(100, 116, 139)

def set_slide_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_header(slide, title_text, category_text="DAWAIFLOW", slide_num="01/12"):
    # Header container
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.6))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    
    # Category prefix
    run_cat = p.add_run()
    run_cat.text = f"● {category_text.upper()}   |   "
    run_cat.font.size = Pt(10)
    run_cat.font.bold = True
    run_cat.font.color.rgb = COLOR_EMERALD
    
    # Hackathon badge
    run_sub = p.add_run()
    run_sub.text = "INNOVATE WITHOUT BORDERS   ·   HEALTHCARE & MEDTECH"
    run_sub.font.size = Pt(9.5)
    run_sub.font.color.rgb = COLOR_MUTED

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.733), Inches(0.8))
    tf2 = title_box.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_WHITE

def add_speaker_notes(slide, script_text, defense_text=""):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    full_text = f"SPEAKER SCRIPT:\n{script_text}\n"
    if defense_text:
        full_text += f"\nJUDGE DEFENSE:\n{defense_text}"
    tf.text = full_text

def create_card(slide, left, top, width, height, bg_color=COLOR_CARD, border_color=COLOR_CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1)
    return shape

def main():
    prs = Presentation()
    # Set 16:9 widescreen dimensions (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Paths to real images
    img_web_dash = "public_site/assets/product/expiryguard-web-app-original.jpg"
    img_mobile_app = "public_site/assets/product/expiryguard-mobile-app-original.jpg"
    img_challan = "uploads/documents/doc_04de78fb23.jpeg"

    # ==========================================
    # SLIDE 1: TITLE & COVER
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, COLOR_BG)
    
    # Title Box
    t_box = s1.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.333), Inches(3.8))
    tf = t_box.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.CENTER
    r0 = p0.add_run()
    r0.text = "DAWAIFLOW"
    r0.font.size = Pt(56)
    r0.font.bold = True
    r0.font.color.rgb = COLOR_WHITE
    
    p1 = tf.add_paragraph()
    p1.alignment = PP_ALIGN.CENTER
    p1.space_before = Pt(14)
    r1 = p1.add_run()
    r1.text = "Making everyday pharmacy work "
    r1.font.size = Pt(22)
    r1.font.color.rgb = COLOR_WHITE
    
    r1_green = p1.add_run()
    r1_green.text = "faster, simpler and smarter."
    r1_green.font.size = Pt(22)
    r1_green.font.bold = True
    r1_green.font.color.rgb = COLOR_EMERALD_LIGHT

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(24)
    r2 = p2.add_run()
    r2.text = "Healthcare & MedTech  |  AI/ML  |  Web & Mobile POS  |  Built for Indian Retail Pharmacies"
    r2.font.size = Pt(13)
    r2.font.color.rgb = COLOR_SECONDARY

    # Connected loop ticker
    flow_box = s1.shapes.add_textbox(Inches(1.5), Inches(5.8), Inches(10.333), Inches(0.8))
    tf_flow = flow_box.text_frame
    p_flow = tf_flow.paragraphs[0]
    p_flow.alignment = PP_ALIGN.CENTER
    r_flow = p_flow.add_run()
    r_flow.text = "PURCHASE   →   INVENTORY   →   BILLING   →   RETURNS   →   BATCH & EXPIRY   →   INSIGHTS"
    r_flow.font.size = Pt(11)
    r_flow.font.bold = True
    r_flow.font.color.rgb = COLOR_MUTED

    add_speaker_notes(
        s1,
        "Good morning esteemed judges. We are Team Dawaiflow, competing under the Innovate Without Borders theme. Dawaiflow is an AI-powered pharmacy management solution built specifically for the reality of Indian retail pharmacies. Our central idea is simple: make everyday pharmacy work faster, simpler, and smarter.",
        "Grounded in actual project code. Fully functional FastAPI backend, Flutter companion app, and native web POS."
    )

    # ==========================================
    # SLIDE 2: THE EVERYDAY PROBLEMS IN A PHARMACY
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, COLOR_BG)
    add_header(s2, "The Everyday Problems in a Pharmacy", "The Reality at the Counter", "02/12")

    # Subtitle
    sub_box = s2.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(11.733), Inches(0.45))
    sub_box.text_frame.paragraphs[0].text = "Existing pharmacy software helps digitise records, but routine operations still leave staff handling repetitive manual work."
    sub_box.text_frame.paragraphs[0].font.size = Pt(13)
    sub_box.text_frame.paragraphs[0].font.color.rgb = COLOR_SECONDARY

    # Left Column: Problem 01 (Billing Manual Steps & Visual Flow)
    create_card(s2, Inches(0.8), Inches(2.15), Inches(5.75), Inches(4.3))
    tb_left = s2.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(5.35), Inches(3.95))
    tf_l = tb_left.text_frame
    tf_l.word_wrap = True

    # 01 Title
    p_l1 = tf_l.paragraphs[0]
    r_tag = p_l1.add_run()
    r_tag.text = "01 · "
    r_tag.font.size = Pt(14)
    r_tag.font.bold = True
    r_tag.font.color.rgb = COLOR_ROSE
    r_head = p_l1.add_run()
    r_head.text = "Billing Still Involves Many Manual Steps"
    r_head.font.size = Pt(14)
    r_head.font.bold = True
    r_head.font.color.rgb = COLOR_WHITE

    # Desc
    p_ld = tf_l.add_paragraph()
    p_ld.space_before = Pt(6)
    p_ld.text = "A pharmacist must search for medicines, select them, enter quantities, check batch details, and complete the bill manually. Small steps repeated for every customer become hours of counter delay."
    p_ld.font.size = Pt(11.5)
    p_ld.font.color.rgb = COLOR_SECONDARY

    # Observation callout
    p_lo = tf_l.add_paragraph()
    p_lo.space_before = Pt(10)
    r_lo_icon = p_lo.add_run()
    r_lo_icon.text = "⏱️ Counter Observation: "
    r_lo_icon.font.bold = True
    r_lo_icon.font.size = Pt(11.5)
    r_lo_icon.font.color.rgb = COLOR_AMBER
    r_lo_text = p_lo.add_run()
    r_lo_text.text = "A routine manual billing interaction takes ~60–65 seconds per customer."
    r_lo_text.font.size = Pt(11.5)
    r_lo_text.font.color.rgb = COLOR_WHITE

    # Step-by-step current flow
    p_flow_title = tf_l.add_paragraph()
    p_flow_title.space_before = Pt(12)
    p_flow_title.text = "TYPICAL CURRENT WORKFLOW:"
    p_flow_title.font.size = Pt(9.5)
    p_flow_title.font.bold = True
    p_flow_title.font.color.rgb = COLOR_MUTED

    p_flow_chain = tf_l.add_paragraph()
    p_flow_chain.space_before = Pt(3)
    p_flow_chain.text = "Customer Waiting  →  Search  →  Select  →  Enter Qty  →  Check  →  Bill  →  Print"
    p_flow_chain.font.size = Pt(10.5)
    p_flow_chain.font.bold = True
    p_flow_chain.font.color.rgb = RGBColor(203, 213, 225)

    # Subtle transition toward Dawaiflow
    p_trans = tf_l.add_paragraph()
    p_trans.space_before = Pt(14)
    r_tr_label = p_trans.add_run()
    r_tr_label.text = "DAWAIFLOW APPROACH:  "
    r_tr_label.font.size = Pt(10)
    r_tr_label.font.bold = True
    r_tr_label.font.color.rgb = COLOR_EMERALD_LIGHT
    r_tr_flow = p_trans.add_run()
    r_tr_flow.text = "SCAN  →  REVIEW  →  BILL"
    r_tr_flow.font.size = Pt(11)
    r_tr_flow.font.bold = True
    r_tr_flow.font.color.rgb = COLOR_WHITE

    # Right Column: Problems 02, 03, 04 (3 stacked cards)
    right_cards = [
        ("02", "Purchase Bills Take Time to Enter",
         "Suppliers provide invoices containing dozens of medicine details. Typing batch codes, expiry dates, and rates manually into the system is repetitive and time-consuming.",
         "Supplier Bill (20–50 rows)  →  Manual Keyboard Entry",
         COLOR_SKY),
        ("03", "Different Tasks Are Not Always Connected",
         "Billing, purchases, stock, returns, and expiry are all connected in a real pharmacy. However, staff often have to manage them as separate, disconnected tasks.",
         "BILLING  |  PURCHASE  |  STOCK  |  RETURNS  |  EXPIRY",
         COLOR_AMBER),
        ("04", "Batch & Expiry Information Needs Constant Attention",
         "Pharmacies deal with thousands of medicines, batches, and expiry dates. Keeping track of them while serving customers and managing daily rush is difficult.",
         "Without automatic FEFO at billing, older batches risk expiring on shelves.",
         COLOR_ROSE)
    ]

    for i, (rc_num, rc_title, rc_desc, rc_flow, rc_col) in enumerate(right_cards):
        y = Inches(2.15 + i * 1.48)
        create_card(s2, Inches(6.8), y, Inches(5.733), Inches(1.36))
        tb_r = s2.shapes.add_textbox(Inches(6.98), y + Inches(0.08), Inches(5.373), Inches(1.2))
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True

        p_rt = tf_r.paragraphs[0]
        r_num = p_rt.add_run()
        r_num.text = f"{rc_num} · "
        r_num.font.size = Pt(12)
        r_num.font.bold = True
        r_num.font.color.rgb = rc_col
        r_title = p_rt.add_run()
        r_title.text = rc_title
        r_title.font.size = Pt(12)
        r_title.font.bold = True
        r_title.font.color.rgb = COLOR_WHITE

        p_rd = tf_r.add_paragraph()
        p_rd.space_before = Pt(2)
        p_rd.text = rc_desc
        p_rd.font.size = Pt(10)
        p_rd.font.color.rgb = COLOR_SECONDARY

        p_rf = tf_r.add_paragraph()
        p_rf.space_before = Pt(3)
        p_rf.text = rc_flow
        p_rf.font.size = Pt(9.5)
        p_rf.font.bold = True
        p_rf.font.color.rgb = rc_col

    # Bottom Fair Reality Ribbon
    create_card(s2, Inches(0.8), Inches(6.58), Inches(11.733), Inches(0.52), COLOR_CARD, COLOR_EMERALD)
    tb_bot = s2.shapes.add_textbox(Inches(1.0), Inches(6.62), Inches(11.333), Inches(0.42))
    tf_bot = tb_bot.text_frame
    p_bot = tf_bot.paragraphs[0]
    r_b1 = p_bot.add_run()
    r_b1.text = "Fair Reality: "
    r_b1.font.bold = True
    r_b1.font.size = Pt(10.5)
    r_b1.font.color.rgb = COLOR_EMERALD_LIGHT
    r_b2 = p_bot.add_run()
    r_b2.text = "Existing pharmacy software helps digitise records, but routine operations still leave staff handling repetitive manual work. Dawaiflow simplifies these exact steps."
    r_b2.font.size = Pt(10.5)
    r_b2.font.color.rgb = COLOR_WHITE

    add_speaker_notes(
        s2,
        "Let us look at what actually happens inside an Indian pharmacy. Existing pharmacy software helps digitise the shop, but routine operations still leave staff handling repetitive manual work. First, billing still involves many manual steps: when a customer arrives, the pharmacist has to search for the medicine, select the right formulation, enter the quantity, verify details, and print. In our counter observation, a routine billing interaction takes around 60 to 65 seconds per customer. Repeat that for 100 or 150 customers every day, and small steps become hours of work. Second, purchase bills take time to enter: wholesale challans arrive with 20 to 50 rows of medicines that staff have to manually type. Third, different pharmacy tasks are not always connected: billing, purchases, inventory, returns, and expiry are naturally linked, but staff often manage them as separate chores. And fourth, batch and expiry information needs constant attention while managing a busy counter. Dawaiflow is built to simplify these exact repetitive steps.",
        "Fair and realistic tone: we do not claim existing software is useless. We focus on the exact repetitive manual steps that current software leaves staff doing by hand (searching, typing rows from paper challans, manually checking expiry). The 60–65s stat is clearly framed as an actual counter observation, not an invented industry claim."
    )

    # ==========================================
    # SLIDE 3: WHY THIS MATTERS
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, COLOR_BG)
    add_header(s3, "Why This Matters", "Practical Impact", "03/12")

    sub_box = s3.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(0.5))
    sub_box.text_frame.paragraphs[0].text = "When routine pharmacy work is repetitive and manual, the operational friction compounds quietly every single day."
    sub_box.text_frame.paragraphs[0].font.size = Pt(13.5)
    sub_box.text_frame.paragraphs[0].font.color.rgb = COLOR_SECONDARY

    # Left 3 steps
    steps = [
        ("Step 1: Excessive Repetitive Work", "Hours lost typing the same medicine names, MRPs, and batch numbers at every stage of the day.", COLOR_ROSE),
        ("Step 2: Rushed Counters & Manual Mistakes", "Mistyped batch numbers, incorrect GST tax rates, and frustratingly long customer queues.", COLOR_AMBER),
        ("Step 3: Silent Expiry & Inventory Desync", "Physical shelf stock does not match software. Expired medicines turn into direct dead financial loss.", COLOR_SKY)
    ]
    for i, (s_title, s_desc, s_col) in enumerate(steps):
        y = Inches(2.3 + i * 1.5)
        card = create_card(s3, Inches(0.8), y, Inches(5.8), Inches(1.3))
        tb = s3.shapes.add_textbox(Inches(1.0), y + Inches(0.12), Inches(5.4), Inches(1.05))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = s_title
        p_t.font.size = Pt(13.5)
        p_t.font.bold = True
        p_t.font.color.rgb = s_col
        
        p_d = tf.add_paragraph()
        p_d.space_before = Pt(4)
        p_d.text = s_desc
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = COLOR_WHITE

    # Right Hero Box
    create_card(s3, Inches(6.9), Inches(2.3), Inches(5.633), Inches(4.3), RGBColor(22, 33, 50), COLOR_EMERALD)
    tb_r = s3.shapes.add_textbox(Inches(7.2), Inches(2.6), Inches(5.033), Inches(3.7))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    
    p_icon = tf_r.paragraphs[0]
    p_icon.alignment = PP_ALIGN.CENTER
    p_icon.text = "⏱️  ↔️  💊"
    p_icon.font.size = Pt(32)

    p_quote = tf_r.add_paragraph()
    p_quote.space_before = Pt(14)
    p_quote.alignment = PP_ALIGN.CENTER
    r_q = p_quote.add_run()
    r_q.text = "Pharmacists should spend time with patients, not typing data twice."
    r_q.font.size = Pt(20)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_WHITE

    p_subquote = tf_r.add_paragraph()
    p_subquote.space_before = Pt(12)
    p_subquote.alignment = PP_ALIGN.CENTER
    p_subquote.text = "In an Indian retail pharmacy, the counter is fast-paced. A pharmacist is a trained healthcare professional dispensing critical medications — not a data entry clerk."
    p_subquote.font.size = Pt(12.5)
    p_subquote.font.color.rgb = COLOR_SECONDARY

    add_speaker_notes(
        s3,
        "Why does this matter? Because when everyday operations rely on manual re-entry, the friction compounds. It leads to slow counter queues, mistyped batch numbers, and expired medicines that sit on shelves until they become a total loss. But most importantly: an Indian retail pharmacist is often the first point of contact for healthcare advice in their neighborhood. A pharmacist should be spending their time checking prescriptions, counseling patients on dosages, and dispensing accurately—not acting as a full-time data entry operator. Pharmacies do not need more complex software; they need a simpler way to manage connected everyday work.",
        "No fake percentage statistics are used. The impact is framed directly in operational friction that every pharmacy owner acknowledges."
    )

    # ==========================================
    # SLIDE 4: OUR SOLUTION (PRODUCT REVEAL)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, COLOR_BG)
    add_header(s4, "Meet Dawaiflow", "Product Reveal", "04/12")

    sub_box = s4.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(0.4))
    sub_box.text_frame.paragraphs[0].text = "“One simple system for the everyday work of a pharmacy.” Dawaiflow connects every operation into one live workflow."
    sub_box.text_frame.paragraphs[0].font.size = Pt(13.5)
    sub_box.text_frame.paragraphs[0].font.color.rgb = COLOR_SECONDARY

    # Left: Web screenshot
    if os.path.exists(img_web_dash):
        create_card(s4, Inches(0.8), Inches(2.1), Inches(6.8), Inches(4.7), RGBColor(0,0,0), COLOR_CARD_BORDER)
        s4.shapes.add_picture(img_web_dash, Inches(0.85), Inches(2.15), width=Inches(6.7))

    # Right: 3 Key value props
    right_points = [
        ("1. Connects the Entire Cycle", "Stock added via purchase bill OCR immediately reflects in live inventory and is ready for billing with FEFO prioritization.", COLOR_EMERALD),
        ("2. Multi-Device by Nature", "Use a smartphone camera for mobile billing and package scanning, while the counter desktop handles high-speed POS and thermal printing.", COLOR_SKY),
        ("3. Practical Counter Reality", "Supports loose tablet splitting, temporary held bills for queue management, customer Khata credit ledgers, and CA tax reports.", COLOR_AMBER)
    ]
    for i, (rp_title, rp_desc, rp_col) in enumerate(right_points):
        y = Inches(2.1 + i * 1.55)
        create_card(s4, Inches(7.8), y, Inches(4.733), Inches(1.4))
        tb = s4.shapes.add_textbox(Inches(8.0), y + Inches(0.12), Inches(4.333), Inches(1.15))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = rp_title
        p_t.font.size = Pt(13.5)
        p_t.font.bold = True
        p_t.font.color.rgb = rp_col
        
        p_d = tf.add_paragraph()
        p_d.space_before = Pt(4)
        p_d.text = rp_desc
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = COLOR_SECONDARY

    add_speaker_notes(
        s4,
        "This is Dawaiflow. Dawaiflow is not just a billing app, and it is not just an inventory spreadsheet. It is one simple, connected system for the everyday work of a pharmacy. When stock arrives, it enters the system; when a customer buys, stock is deducted; when batches approach expiry, the system prioritizes them; and when returns occur, the original batch is restored automatically. On screen right now, you are looking at our actual web dashboard from our active pilot. It is tracking over 10,000 active inventory items, 182 daily bills, and ₹25,000 in sales, with real-time sync between the mobile app and the counter terminal. Everything is connected in one place.",
        "Show the real screenshot on the slide. Point out the 'Live Sync with Mobile' badge and the 10,000 active items from our actual database."
    )

    # ==========================================
    # SLIDE 5: HOW IT WORKS
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, COLOR_BG)
    add_header(s5, "How Dawaiflow Works at the Counter", "Real-World Flow", "05/12")

    sub_box = s5.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(0.4))
    sub_box.text_frame.paragraphs[0].text = "From incoming stock to counter checkout — here is how data moves through Dawaiflow without re-typing."
    sub_box.text_frame.paragraphs[0].font.size = Pt(13.5)
    sub_box.text_frame.paragraphs[0].font.color.rgb = COLOR_SECONDARY

    steps_hw = [
        ("STEP 01", "Intake or Scan", "Pharmacist uploads a photo of a distributor challan, or points their phone camera at medicine strips on the counter.", "✓ Reads 1–7 strips at once\n✓ Reads paper challans\n✓ Local barcode pre-check", COLOR_EMERALD),
        ("STEP 02", "Intelligent Extraction & FEFO", "AI & computer vision extract medicine names, batches, expiry, and MRP. System automatically assigns the earliest-expiring batch (FEFO).", "✓ Brand vs Salt distinction\n✓ First-Expiring, First-Out (FEFO)\n✓ Loose-tablet price math", COLOR_SKY),
        ("STEP 03", "Instant Bill & Live Sync", "Tax invoice is generated in 1 click. Stock is deducted, counter printer prints receipt in <100ms, and web dashboard updates live.", "✓ GST-compliant invoice\n✓ ESC/POS thermal print\n✓ Live mobile ↔ desktop sync", COLOR_AMBER)
    ]
    for i, (s_num, s_title, s_desc, s_bullets, s_col) in enumerate(steps_hw):
        x = Inches(0.8 + i * 4.0)
        create_card(s5, x, Inches(2.2), Inches(3.733), Inches(4.0))
        tb = s5.shapes.add_textbox(x + Inches(0.2), Inches(2.4), Inches(3.333), Inches(3.6))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_num = tf.paragraphs[0]
        p_num.text = s_num
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = s_col
        
        p_title = tf.add_paragraph()
        p_title.space_before = Pt(4)
        p_title.text = s_title
        p_title.font.size = Pt(15)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_WHITE

        p_desc = tf.add_paragraph()
        p_desc.space_before = Pt(8)
        p_desc.text = s_desc
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = COLOR_SECONDARY

        p_bul = tf.add_paragraph()
        p_bul.space_before = Pt(14)
        p_bul.text = s_bullets
        p_bul.font.size = Pt(11)
        p_bul.font.color.rgb = s_col

    # Bottom summary tag
    create_card(s5, Inches(0.8), Inches(6.35), Inches(11.733), Inches(0.6), COLOR_CARD, COLOR_EMERALD)
    tb_b = s5.shapes.add_textbox(Inches(1.0), Inches(6.4), Inches(11.333), Inches(0.5))
    tb_b.text_frame.paragraphs[0].text = "Zero Re-entry: One single scan completes billing, inventory deduction, batch tracking, and accounting simultaneously."
    tb_b.text_frame.paragraphs[0].font.size = Pt(12.5)
    tb_b.text_frame.paragraphs[0].font.bold = True
    tb_b.text_frame.paragraphs[0].font.color.rgb = COLOR_WHITE

    add_speaker_notes(
        s5,
        "Here is how Dawaiflow works in practice. In Step 1, when wholesale stock arrives or a customer places medicines on the counter, the pharmacist takes a photo. Our system can read up to 7 medicine strips in a single photo, or process an entire 30-item paper challan. In Step 2, our vision engine extracts the medicine name, strength, batch number, and expiry date. Crucially, the system automatically allocates the earliest expiring batch using FEFO—First-Expiring, First-Out. In Step 3, the bill is created with one click, stock is deducted across all devices, and our desktop agent triggers the thermal receipt printer in under 100 milliseconds. One action updates billing, inventory, batch records, and taxes simultaneously.",
        "Explain our Win32 print agent (dawaiflow_print_agent.py) that connects to actual 58mm/80mm ESC/POS counter printers via WebSockets."
    )

    # ==========================================
    # SLIDE 6: WHAT MAKES DAWAIFLOW DIFFERENT?
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, COLOR_BG)
    add_header(s6, "What Makes Dawaiflow Different?", "Product Philosophy", "06/12")

    sub_box = s6.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(0.4))
    sub_box.text_frame.paragraphs[0].text = "“Dawaiflow does not just manage a pharmacy. It connects the work happening inside it.”"
    sub_box.text_frame.paragraphs[0].font.size = Pt(13.5)
    sub_box.text_frame.paragraphs[0].font.color.rgb = COLOR_SECONDARY

    diffs = [
        ("01", "AI Where It Actually Helps", "Not AI for hype. AI is strictly deployed to eliminate repetitive keyboard work: reading complex medicine packaging, batch numbers, and messy paper challans.", COLOR_EMERALD),
        ("02", "One Connected Data Loop", "Purchase intake, live inventory, counter POS billing, customer returns, and expiry tracking are not separate modules. They feed the exact same live ledger.", COLOR_SKY),
        ("03", "Built for Indian Counter Reality", "Understands loose tablet strip math, holds parked bills when customers step aside to pick another item, and imports historical Marg ERP / Excel databases seamlessly.", COLOR_AMBER),
        ("04", "Action Over Passive Reports", "Traditional software shows an expiry report after the medicine is already dead. Dawaiflow enforces FEFO at billing time, selling expiring batches first to prevent waste.", COLOR_ROSE)
    ]

    for i, (d_num, d_title, d_desc, d_col) in enumerate(diffs):
        col_idx = i % 2
        row_idx = i // 2
        x = Inches(0.8 + col_idx * 5.95)
        y = Inches(2.2 + row_idx * 2.3)
        create_card(s6, x, y, Inches(5.75), Inches(2.05))
        tb = s6.shapes.add_textbox(x + Inches(0.2), y + Inches(0.15), Inches(5.35), Inches(1.75))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_t = tf.paragraphs[0]
        r_num = p_t.add_run()
        r_num.text = f"{d_num}  ·  "
        r_num.font.size = Pt(14)
        r_num.font.bold = True
        r_num.font.color.rgb = d_col
        
        r_title = p_t.add_run()
        r_title.text = d_title
        r_title.font.size = Pt(14)
        r_title.font.bold = True
        r_title.font.color.rgb = COLOR_WHITE

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(8)
        p_d.text = d_desc
        p_d.font.size = Pt(11.8)
        p_d.font.color.rgb = COLOR_SECONDARY

    add_speaker_notes(
        s6,
        "What makes Dawaiflow truly different from existing pharmacy software? First, our AI is practical, not decorative. We do not use AI for marketing buzzwords; we use it strictly to eliminate the manual typing of medicine names, batch numbers, and distributor invoices. Second, Dawaiflow is one connected data loop. In traditional software, billing and purchase entry are separated, creating discrepancies. Here, they feed the exact same live ledger. Third, Dawaiflow is built for Indian counter realities: it calculates loose tablet prices when a customer wants only 2 pills from a strip of 10, and it allows staff to park a bill with one tap when a customer steps aside to pick another item. And fourth, Dawaiflow drives action: instead of showing a sad report of expired medicines at the end of the month, it enforces FEFO at billing time so expiring medicines sell first.",
        "Highlight our HeldBill model and tablets_per_strip loose tablet database architecture in models.py."
    )

    # ==========================================
    # SLIDE 7: AI IN DAWAIFLOW
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, COLOR_BG)
    add_header(s7, "AI That Actually Helps", "Practical Intelligence", "07/12")

    sub_box = s7.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(0.4))
    sub_box.text_frame.paragraphs[0].text = "AI handles repetitive reading and data extraction so pharmacy staff can focus on the work that matters."
    sub_box.text_frame.paragraphs[0].font.size = Pt(13.5)
    sub_box.text_frame.paragraphs[0].font.color.rgb = COLOR_SECONDARY

    # Left: Real Challan Photo
    if os.path.exists(img_challan):
        create_card(s7, Inches(0.8), Inches(2.1), Inches(4.5), Inches(4.7), RGBColor(0,0,0), COLOR_CARD_BORDER)
        s7.shapes.add_picture(img_challan, Inches(0.85), Inches(2.15), width=Inches(4.4))

    # Right: 3 Real AI capabilities
    ai_points = [
        ("1. Distributor Challan & Invoice OCR", "Converts complex, multi-item printed & handwritten supplier invoices (like Marg challans) directly into structured inventory stock batches in seconds.", "Gemini Vision", COLOR_EMERALD),
        ("2. Multi-Strip Counter Billing Scan", "A single smartphone photo captures 1 to 7 distinct medicine packages on the counter. AI distinguishes brand names from chemical salts, batch numbers, and expiry.", "Multi-Item Vision", COLOR_SKY),
        ("3. Local Edge Pre-Detection", "Barcodes and QR codes are detected locally in <20ms using OpenCV and PaddleOCR before calling cloud vision, saving API costs and providing instant billing feedback.", "OpenCV + PaddleOCR", COLOR_AMBER)
    ]

    for i, (ap_title, ap_desc, ap_tag, ap_col) in enumerate(ai_points):
        y = Inches(2.1 + i * 1.55)
        create_card(s7, Inches(5.6), y, Inches(6.933), Inches(1.4))
        tb = s7.shapes.add_textbox(Inches(5.8), y + Inches(0.12), Inches(6.533), Inches(1.15))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_t = tf.paragraphs[0]
        r_t = p_t.add_run()
        r_t.text = ap_title
        r_t.font.size = Pt(13.5)
        r_t.font.bold = True
        r_t.font.color.rgb = ap_col
        
        r_tag = p_t.add_run()
        r_tag.text = f"   [{ap_tag}]"
        r_tag.font.size = Pt(10)
        r_tag.font.color.rgb = COLOR_MUTED

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(4)
        p_d.text = ap_desc
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = COLOR_SECONDARY

    add_speaker_notes(
        s7,
        "Let us talk about AI honestly. The photo on the left is a real distributor challan from Gupta Enterprises that we tested in our project. Notice that it contains printed items from Marg ERP mixed with handwritten medicine names and batch numbers at the bottom. This is what Indian pharmacy owners deal with every day. Dawaiflow’s AI pipeline extracts these rows directly into structured inventory batches in seconds. For counter billing, our multi-item vision service can identify up to 7 distinct medicine strips in one photograph, correctly separating the commercial brand name from the chemical formulation. And to ensure lightning speed and cost-effectiveness, we built an edge pre-check with OpenCV and PaddleOCR that scans barcodes in under 20 milliseconds before any cloud API is called. AI is used where it genuinely saves human time.",
        "Point to ai/multi_item_scan_service.py and ai/ocr_service.py. This dual-tier architecture demonstrates engineering maturity and cost awareness."
    )

    # ==========================================
    # SLIDE 8: IMPACT & BENEFITS (BEFORE / AFTER)
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, COLOR_BG)
    add_header(s8, "What Changes for a Pharmacy?", "Operational Impact", "08/12")

    sub_box = s8.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(0.4))
    sub_box.text_frame.paragraphs[0].text = "Practical improvements measured directly in saved minutes, fewer errors, and protected inventory value."
    sub_box.text_frame.paragraphs[0].font.size = Pt(13.5)
    sub_box.text_frame.paragraphs[0].font.color.rgb = COLOR_SECONDARY

    # Native Table Shape
    rows = [
        ("Receiving Stock", "20–30 minutes typing rows of medicines, batch codes, rates, and expiry dates manually.", "Snap 1 photo of the challan. AI populates verified inventory batches in seconds."),
        ("Counter Billing", "Customer waits while staff writes handwritten slips or slowly searches names in old software.", "Quick camera scan or instant 2-keystroke lookup. Bill generated & auto-printed in seconds."),
        ("Expiry Management", "Periodic manual checking of shelves. Expired medicines discovered too late.", "FEFO billing automatically allocates older batches first; 60-day expiry radar alerts staff."),
        ("Customer Returns", "Manual paper math, forgotten batch adjustments, inaccurate inventory counts.", "1-click return restores stock to the original batch and logs refund atomically."),
        ("Daily Closing", "Hours reconciling registers, cash tallies, and preparing GST sheets for accountant.", "Instant dashboard totals, multi-mode payment breakdown, and 1-click CA tax report share.")
    ]

    table_shape = s8.shapes.add_table(6, 3, Inches(0.8), Inches(2.2), Inches(11.733), Inches(4.2))
    table = table_shape.table
    table.columns[0].width = Inches(2.5)
    table.columns[1].width = Inches(4.616)
    table.columns[2].width = Inches(4.616)

    # Headers
    headers = ["PHARMACY TASK", "BEFORE DAWAIFLOW (MANUAL)", "WITH DAWAIFLOW (CONNECTED)"]
    h_colors = [COLOR_WHITE, COLOR_ROSE, COLOR_EMERALD_LIGHT]
    for col_idx in range(3):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD
        p = cell.text_frame.paragraphs[0]
        p.text = headers[col_idx]
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = h_colors[col_idx]

    # Data Rows
    for r_idx, row_data in enumerate(rows):
        for c_idx in range(3):
            cell = table.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(22, 30, 46) if r_idx % 2 == 0 else COLOR_CARD
            p = cell.text_frame.paragraphs[0]
            p.text = row_data[c_idx]
            p.font.size = Pt(10.5)
            if c_idx == 0:
                p.font.bold = True
                p.font.color.rgb = COLOR_WHITE
            elif c_idx == 1:
                p.font.color.rgb = COLOR_SECONDARY
            else:
                p.font.bold = True
                p.font.color.rgb = COLOR_EMERALD_LIGHT

    add_speaker_notes(
        s8,
        "What actually changes when a pharmacy adopts Dawaiflow? This comparison table shows the direct operational shift. When receiving stock, staff go from typing 30 minutes of spreadsheet rows to snapping one photo of the bill. Counter billing moves from slow manual searching to an instant camera scan or two keystrokes. Expiry management shifts from periodic shelf-digging to automatic FEFO batch allocation at the point of sale. Customer returns restore stock to the exact original batch with one click. And at closing time, the pharmacy owner does not spend hours calculating tallies—they have an instant financial summary and a one-click GST export ready for their chartered accountant.",
        "Emphasize our CaProfile and CaShareLog models which allow pharmacies to share GSTR reports directly with their accountants."
    )

    # ==========================================
    # SLIDE 9: TECHNOLOGY
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, COLOR_BG)
    add_header(s9, "How We Built Dawaiflow", "Technical Architecture", "09/12")

    sub_box = s9.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(0.4))
    sub_box.text_frame.paragraphs[0].text = "Built with high-performance, proven technologies designed for low latency, offline resiliency, and hardware compatibility."
    sub_box.text_frame.paragraphs[0].font.size = Pt(13.5)
    sub_box.text_frame.paragraphs[0].font.color.rgb = COLOR_SECONDARY

    tech_pillars = [
        ("FRONTEND & POS", "Web & Mobile", "• Web: Native ES6, CSS3, Inter + JetBrains Mono (Zero bloated JS dependencies).\n• Mobile: Flutter (Dart) Android companion app for counter camera billing.", COLOR_SKY),
        ("BACKEND & API", "FastAPI & Python", "• Python 3.11 with FastAPI.\n• Asynchronous REST API + real-time WebSockets for sub-100ms hardware event streaming.\n• Background workers for scheduled alerts.", COLOR_EMERALD),
        ("DATA ENGINE", "Postgres + 240k Master", "• PostgreSQL with SQLAlchemy.\n• Pre-indexed Indian Medicine Catalog of 240,000+ medicines with composition & HSN.\n• Atomic transactions & soft-delete.", COLOR_AMBER),
        ("AI & HARDWARE", "Vision & Print Agent", "• AI: Gemini 3.1 Flash-Lite + PaddleOCR + OpenCV.\n• Hardware: Windows Desktop Print Agent piping ESC/POS to 58mm/80mm thermal printers.", COLOR_ROSE)
    ]

    for i, (tp_cat, tp_title, tp_desc, tp_col) in enumerate(tech_pillars):
        x = Inches(0.8 + i * 2.98)
        create_card(s9, x, Inches(2.2), Inches(2.78), Inches(3.9))
        tb = s9.shapes.add_textbox(x + Inches(0.15), Inches(2.35), Inches(2.48), Inches(3.6))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_cat = tf.paragraphs[0]
        p_cat.text = tp_cat
        p_cat.font.size = Pt(10.5)
        p_cat.font.bold = True
        p_cat.font.color.rgb = tp_col
        
        p_t = tf.add_paragraph()
        p_t.space_before = Pt(4)
        p_t.text = tp_title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(10)
        p_d.text = tp_desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_SECONDARY

    # Bottom flow banner
    create_card(s9, Inches(0.8), Inches(6.3), Inches(11.733), Inches(0.6), COLOR_CARD, COLOR_SKY)
    tb_arch = s9.shapes.add_textbox(Inches(1.0), Inches(6.35), Inches(11.333), Inches(0.5))
    tb_arch.text_frame.paragraphs[0].text = "USER DEVICE (Phone / PC)   ⇄   FastAPI (Async REST & WebSockets)   ⇄   PostgreSQL (240k Catalog)   ⇄   AI Vision   ⇄   ESC/POS Thermal Spooler"
    tb_arch.text_frame.paragraphs[0].font.size = Pt(10.5)
    tb_arch.text_frame.paragraphs[0].font.bold = True
    tb_arch.text_frame.paragraphs[0].font.color.rgb = COLOR_WHITE

    add_speaker_notes(
        s9,
        "Behind Dawaiflow is a production-grade, highly reliable technical stack. The web POS is built with native modern ES6 and CSS3, ensuring zero JavaScript framework overhead and instant page loads. Our mobile app is built in Flutter. The backend is written in Python 3.11 with FastAPI, providing asynchronous REST endpoints and real-time WebSockets for live hardware communication. Our database is PostgreSQL, pre-indexed with a master catalog of over 240,000 Indian medicines with their composition and HSN codes. On the hardware side, our dedicated Windows Print Agent communicates directly with counter thermal printers using raw ESC/POS spooling. It is designed to work reliably in low-bandwidth, high-volume counter environments.",
        "Cite indian_medicine_data.csv (240k records) and the lightweight dependency footprint in requirements.txt."
    )

    # ==========================================
    # SLIDE 10: FEASIBILITY & IMPLEMENTATION
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, COLOR_BG)
    add_header(s10, "Built to Work in the Real World", "Feasibility & Pilot", "10/12")

    sub_box = s10.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(0.4))
    sub_box.text_frame.paragraphs[0].text = "Dawaiflow is not an abstract concept. It is built to work immediately on the equipment Indian pharmacies already possess."
    sub_box.text_frame.paragraphs[0].font.size = Pt(13.5)
    sub_box.text_frame.paragraphs[0].font.color.rgb = COLOR_SECONDARY

    # Left: Mobile Pilot Screenshot
    if os.path.exists(img_mobile_app):
        create_card(s10, Inches(0.8), Inches(2.1), Inches(3.8), Inches(4.7), RGBColor(0,0,0), COLOR_CARD_BORDER)
        s10.shapes.add_picture(img_mobile_app, Inches(0.85), Inches(2.15), width=Inches(3.7))

    # Right: 4 Feasibility Pillars
    feasibility_points = [
        ("1. Zero Hardware Cost", "Works on standard Android smartphones and existing counter PCs. No expensive barcode guns or proprietary scanners needed.", COLOR_EMERALD),
        ("2. Marg ERP & Excel Coexistence", "Includes direct spreadsheet and Marg ERP import tools. Pharmacies can onboard their entire 10,000-item inventory in minutes without disruption.", COLOR_SKY),
        ("3. Real Pilot Validation", "Actively tested with Indian pharmacies (e.g. Vashist Pharmacy) processing 180+ bills/day, verifying live counter speed and FEFO alerts.", COLOR_AMBER),
        ("4. Resilient & Fail-Safe", "In-memory preloading for instant 0ms billing UI, dual WebSocket/HTTP print spooler, and strict transactional idempotency.", COLOR_ROSE)
    ]

    for i, (fp_title, fp_desc, fp_col) in enumerate(feasibility_points):
        y = Inches(2.1 + i * 1.18)
        create_card(s10, Inches(4.9), y, Inches(7.633), Inches(1.05))
        tb = s10.shapes.add_textbox(Inches(5.1), y + Inches(0.08), Inches(7.233), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_t = tf.paragraphs[0]
        p_t.text = fp_title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = fp_col

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(3)
        p_d.text = fp_desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_SECONDARY

    add_speaker_notes(
        s10,
        "Dawaiflow was designed from day one to be feasible in the real world. First, it requires zero investment in new hardware: pharmacy staff use the Android smartphones they already own and the Windows PCs they already have at their counters. Second, we do not ask pharmacies to abandon their past data. We built direct Excel and Marg ERP data migration tools so a pharmacy can import their 10,000 inventory items in minutes. Third, this is already proven in an active pilot at Vashist Pharmacy, processing over 180 customer bills daily. And fourth, the system is fail-safe: with in-memory caching and dual WebSocket-plus-HTTP print streaming, the counter never freezes during peak hours.",
        "Point out the live pilot lead model (PilotLead) and the authentic photo of the app running on an actual smartphone on a pharmacy counter."
    )

    # ==========================================
    # SLIDE 11: SCALABILITY
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, COLOR_BG)
    add_header(s11, "Where Dawaiflow Can Go Next", "Scalability & Roadmap", "11/12")

    sub_box = s11.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.733), Inches(0.4))
    sub_box.text_frame.paragraphs[0].text = "“Start with one pharmacy. Scale the same connected workflow across many.”"
    sub_box.text_frame.paragraphs[0].font.size = Pt(13.5)
    sub_box.text_frame.paragraphs[0].font.color.rgb = COLOR_SECONDARY

    stages = [
        ("STAGE 1 (ACTIVE)", "Single Counter", "Complete connected workflow: AI invoice intake, live FEFO inventory, rapid counter billing, and thermal print spooling for independent pharmacies.", COLOR_EMERALD),
        ("STAGE 2 (BUILT IN CODE)", "Multi-Branch Stores", "Centralized stock visibility across multiple retail outlets. Transfer stock between branches before it expires; unified owner dashboard.", COLOR_SKY),
        ("STAGE 3", "Distributor Link", "Automated reordering based on actual counter sales velocity and supplier payment terms. 1-click PO generation directly to distributors.", COLOR_AMBER),
        ("STAGE 4", "Network Intelligence", "City-level medicine availability and shortage detection. Helping patients find scarce medicines across connected pharmacy networks.", COLOR_ROSE)
    ]

    for i, (st_stage, st_title, st_desc, st_col) in enumerate(stages):
        x = Inches(0.8 + i * 2.98)
        create_card(s11, x, Inches(2.3), Inches(2.78), Inches(3.7))
        tb = s11.shapes.add_textbox(x + Inches(0.15), Inches(2.45), Inches(2.48), Inches(3.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_st = tf.paragraphs[0]
        p_st.text = st_stage
        p_st.font.size = Pt(10.5)
        p_st.font.bold = True
        p_st.font.color.rgb = st_col
        
        p_t = tf.add_paragraph()
        p_t.space_before = Pt(4)
        p_t.text = st_title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(10)
        p_d.text = st_desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_SECONDARY

    create_card(s11, Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.65), COLOR_CARD, COLOR_EMERALD)
    tb_sc = s11.shapes.add_textbox(Inches(1.0), Inches(6.32), Inches(11.333), Inches(0.5))
    tb_sc.text_frame.paragraphs[0].text = "Every single expansion stage relies on the exact same core truth: clean, connected data captured effortlessly at the counter."
    tb_sc.text_frame.paragraphs[0].font.size = Pt(12)
    tb_sc.text_frame.paragraphs[0].font.bold = True
    tb_sc.text_frame.paragraphs[0].font.color.rgb = COLOR_WHITE

    add_speaker_notes(
        s11,
        "Where does Dawaiflow go next? Our scalability roadmap is grounded and modular. We have built and validated Stage 1 for independent pharmacies. Stage 2 is multi-branch store networks—and our database schema already includes multi-branch models and stock transfer tracking. If one store has an excess batch expiring in 60 days, it can be transferred to a high-volume branch to prevent waste. Stage 3 is automated distributor reordering based on live sales velocity. And Stage 4 is regional network intelligence, enabling neighboring pharmacies to coordinate during medicine shortages. Every single stage builds upon the same core foundation: effortless, accurate data capture at the counter.",
        "Show that StoreBranch is already coded into models.py and web/branches.html."
    )

    # ==========================================
    # SLIDE 12: CLOSING
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12, COLOR_BG)

    tb_close = s12.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5))
    tf_c = tb_close.text_frame
    tf_c.word_wrap = True
    
    p_c0 = tf_c.paragraphs[0]
    p_c0.alignment = PP_ALIGN.CENTER
    r_c0 = p_c0.add_run()
    r_c0.text = "● CONCLUSION"
    r_c0.font.size = Pt(12)
    r_c0.font.bold = True
    r_c0.font.color.rgb = COLOR_EMERALD

    p_c1 = tf_c.add_paragraph()
    p_c1.space_before = Pt(14)
    p_c1.alignment = PP_ALIGN.CENTER
    r_c1 = p_c1.add_run()
    r_c1.text = "A Smarter Way to Run a Pharmacy"
    r_c1.font.size = Pt(38)
    r_c1.font.bold = True
    r_c1.font.color.rgb = COLOR_WHITE

    p_c2 = tf_c.add_paragraph()
    p_c2.space_before = Pt(14)
    p_c2.alignment = PP_ALIGN.CENTER
    r_c2 = p_c2.add_run()
    r_c2.text = "“Less manual work. More connected pharmacy operations.”"
    r_c2.font.size = Pt(22)
    r_c2.font.bold = True
    r_c2.font.color.rgb = COLOR_EMERALD_LIGHT

    p_c3 = tf_c.add_paragraph()
    p_c3.space_before = Pt(28)
    p_c3.alignment = PP_ALIGN.CENTER
    r_c3 = p_c3.add_run()
    r_c3.text = "DAWAIFLOW\nMaking everyday pharmacy work faster, simpler and smarter."
    r_c3.font.size = Pt(15)
    r_c3.font.color.rgb = COLOR_SECONDARY

    p_c4 = tf_c.add_paragraph()
    p_c4.space_before = Pt(24)
    p_c4.alignment = PP_ALIGN.CENTER
    r_c4 = p_c4.add_run()
    r_c4.text = "Healthcare & MedTech   ·   Built by Team Dawaiflow   ·   Ready for Live Demo"
    r_c4.font.size = Pt(12)
    r_c4.font.color.rgb = COLOR_MUTED

    add_speaker_notes(
        s12,
        "To conclude: running a pharmacy should not mean spending hours typing the same information into disconnected tools. Dawaiflow delivers less manual work and more connected pharmacy operations. It is a real product, solving a real problem, with real AI, validated on real Indian pharmacy counter data. Thank you so much for your time. We invite you to test our live billing scan, see our thermal printing in action, and ask us any technical or business questions.",
        "Invite judges to test a live bill creation, multi-item strip scan, or thermal print dispatch."
    )

    # Save presentation
    try:
        prs.save(OUTPUT_FILE)
        print(f"Presentation saved successfully to: {OUTPUT_FILE}")
    except PermissionError:
        fallback_file = "Dawaiflow_Hackathon_Presentation_Updated.pptx"
        prs.save(fallback_file)
        print(f"Target file was locked by PowerPoint. Saved successfully to: {fallback_file}")

if __name__ == "__main__":
    main()
